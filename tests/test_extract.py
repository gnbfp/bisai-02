"""M1 抽取管线的单测（requirements §7.5，D-15 / D-21 / D-25）。

夹具在 tmp_path 里现场生成：仓库不提交二进制测试文件（tests/fixtures 保持为空）。
生成 PDF 的文字一律用 ASCII —— 内置字体没有中文字形，中文会抽出乱码，
这是测试夹具的字体问题，不是抽取器的问题（真实作业书由 PyMuPDF 按原字形抽出）。
"""

import pytest

from src.intelligence.extract import (
    TABLE_MARKER,
    ExtractError,
    check_deadline,
    check_meta_fields,
    check_radical_residue,
    check_weight_sum,
    extract_text,
    normalize_cjk,
)
from src.models import AssignmentMeta, RubricPoint, SchemaError


def _make_pdf(path, body, table):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in body:
        page.insert_text(fitz.Point(72, y), line, fontsize=11)
        y += 16

    col_w, row_h, x0 = 130, 18, 72
    top = y + 12
    n_cols = max(len(row) for row in table)
    for r in range(len(table) + 1):
        yy = top + r * row_h
        page.draw_line(fitz.Point(x0, yy), fitz.Point(x0 + col_w * n_cols, yy))
    for c in range(n_cols + 1):
        xx = x0 + c * col_w
        page.draw_line(fitz.Point(xx, top), fitz.Point(xx, top + len(table) * row_h))
    for r, row in enumerate(table):
        for c, value in enumerate(row):
            page.insert_text(fitz.Point(x0 + c * col_w + 4, top + r * row_h + 13), value, fontsize=10)
    doc.save(str(path))
    doc.close()


def _make_docx(path, paragraphs, table=()):
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        t.style = "Table Grid"
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    doc.save(str(path))


def _point(pid, weight=None):
    return RubricPoint(id=pid, quote=f"{pid} 原文", observable="可核对", weight=weight)


# ---------- 分派与拒收 ----------


def test_text_passthrough(tmp_path):
    p = tmp_path / "brief.txt"
    p.write_text("作业书正文", encoding="utf-8")
    assert extract_text(p) == "作业书正文"


def test_image_is_rejected(tmp_path):
    p = tmp_path / "scan.png"
    p.write_bytes(b"not really an image")
    with pytest.raises(ExtractError) as exc:
        extract_text(p)
    assert "图片" in str(exc.value)


def test_unknown_suffix_is_rejected(tmp_path):
    p = tmp_path / "brief.pptx"
    p.write_bytes(b"x")
    with pytest.raises(ExtractError):
        extract_text(p)


def test_scanned_pdf_is_rejected(tmp_path):
    import fitz

    p = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(p))
    doc.close()
    with pytest.raises(ExtractError) as exc:
        extract_text(p)
    assert "文字层" in str(exc.value)


# ---------- PDF ----------


def test_pdf_body_and_table(tmp_path):
    p = tmp_path / "a.pdf"
    _make_pdf(p, ["Assignment brief"], [["Item", "Score"], ["Feature", "40"], ["Report", "60"]])
    text = extract_text(p)
    assert "Assignment brief" in text
    assert TABLE_MARKER in text
    assert "- Item：Feature｜Score：40" in text
    assert "- Item：Report｜Score：60" in text


def test_pdf_table_only_still_extractable(tmp_path):
    p = tmp_path / "t.pdf"
    _make_pdf(p, [], [["Item", "Score"], ["Feature", "100"]])
    text = extract_text(p)
    assert TABLE_MARKER in text
    assert "- Item：Feature｜Score：100" in text


# ---------- Word ----------


def test_docx_body_and_table(tmp_path):
    p = tmp_path / "a.docx"
    _make_docx(p, ["作业书正文"], [["评分项", "分值"], ["功能实现", "40"], ["报告", "60"]])
    text = extract_text(p)
    assert "作业书正文" in text
    assert TABLE_MARKER in text
    assert "- 评分项：功能实现｜分值：40" in text


def test_docx_empty_cells_are_skipped(tmp_path):
    p = tmp_path / "b.docx"
    _make_docx(p, [], [["评分项", "分值"], ["功能实现", ""]])
    text = extract_text(p)
    assert "- 评分项：功能实现" in text
    assert "分值：" not in text


def test_empty_docx_is_rejected(tmp_path):
    p = tmp_path / "e.docx"
    _make_docx(p, [])
    with pytest.raises(ExtractError):
        extract_text(p)


# ---------- 权重加总软校验（§7.5 / D-25）----------


def test_weight_sum_in_range_is_silent():
    assert check_weight_sum([_point("R1", 40), _point("R2", 60)]) is None
    assert check_weight_sum([_point("R1", 95)]) is None


def test_weight_sum_all_missing_is_skipped():
    assert check_weight_sum([_point("R1")]) is None
    assert check_weight_sum([_point("R1"), _point("R2")]) is None


def test_weight_sum_partial_missing_warns():
    # D-48：一半带分一半不带，往往是 M1 把正文要求也当评分点收进来了
    warning = check_weight_sum([_point("R1", 60), _point("R2")])
    assert warning is not None
    assert "1/2 条没有分值" in warning
    assert "疑似把正文要求当成了评分点" in warning


def test_weight_sum_twenty_scale_is_skipped():
    # 20 分制：能归一化到百分制，但不是百分制量纲 → 整体跳过，不误报
    assert check_weight_sum([_point("R1", 10), _point("R2", 10)]) is None


def test_weight_sum_lost_rows_warns():
    warning = check_weight_sum([_point("R1", 40), _point("R2", 45)])
    assert warning is not None
    assert "丢行" in warning
    assert "85" in warning


def test_weight_sum_over_hundred_warns():
    assert check_weight_sum([_point("R1", 70), _point("R2", 50)]) is not None


def test_weight_sum_of_empty_rubric_is_silent():
    assert check_weight_sum([]) is None


# ---------- 截止时间软校验（D-49）----------


def _meta(deadline=""):
    return AssignmentMeta(
        course="编译原理",
        title="课程设计",
        submission="源码 + 报告",
        deadline=deadline,
        source_file="作业书.txt",
    )


def test_empty_deadline_is_valid_and_warns():
    meta = _meta("")
    meta.validate()                                   # 空 deadline 不再抛 SchemaError（D-49）
    warning = check_deadline(meta)
    assert warning is not None
    assert "未标注" in warning


def test_placeholder_deadline_warns():
    warning = check_deadline(_meta("1970-01-01T00:00"))
    assert warning is not None
    assert "占位值" in warning


def test_real_deadline_is_silent():
    assert check_deadline(_meta("2026-06-30T00:00")) is None


# ---------- 作业元信息缺件软校验（2026-09-17 PM 拍 A）----------


def _meta_fields(course, title, submission):
    return AssignmentMeta(
        course=course,
        title=title,
        submission=submission,
        deadline="2026-06-30T00:00",
        source_file="指导书.docx",
    )


def test_empty_course_is_valid_and_warns():
    """只有格式要求的指导书：填不出的字段允许空，但要软警告点名。"""
    meta = _meta_fields("", "课程设计报告", "源码")
    meta.validate()                          # 空 course 不再抛 SchemaError（拍 A）
    warning = check_meta_fields(meta)
    assert warning is not None
    assert "课程名" in warning
    assert "未标注" in warning


def test_all_missing_meta_fields_are_named():
    warning = check_meta_fields(_meta_fields("", "", ""))
    assert "课程名" in warning and "作业标题" in warning and "提交物" in warning


def test_complete_meta_is_silent():
    assert check_meta_fields(_meta_fields("编译原理", "课程设计", "源码")) is None


def test_source_file_is_still_required():
    """程序给的事实（文件名）不许空 —— 这是唯一保留的必填。"""
    with pytest.raises(SchemaError):
        AssignmentMeta(course="", title="", submission="", deadline="", source_file="").validate()


# ---------- CJK 部首归一化（D-43）----------


def test_normalize_cjk_kangxi_radicals():
    assert normalize_cjk("\u2f00") == "一"
    assert normalize_cjk("\u2f2f") == "工"
    assert normalize_cjk("\u2f6c") == "目"


def test_normalize_cjk_simplified_radicals():
    """NFKC 覆盖不到的，靠显式小表兜住。"""
    assert normalize_cjk("\u2eda") == "页"
    assert normalize_cjk("\u2edb") == "风"
    assert normalize_cjk("\u2ed4") == "门"
    assert normalize_cjk("\u2ee2") == "马"
    assert normalize_cjk("\u2ec6") == "角"


def test_normalize_cjk_maps_jian_radical():
    """⻅ 是 BIM 任务书 91 个部首里 NFKC 兜不住的最后一个。"""
    assert normalize_cjk("详\u2ec5") == "详见"


def test_normalize_cjk_real_sentence():
    """真实文档里抽出来的那种句子。"""
    assert normalize_cjk("作\u2edb") == "作风"
    assert normalize_cjk("入\u2ed4") == "入门"


def test_normalize_cjk_length_unchanged():
    """逐字符映射 => 长度必须不变（引用定位、M8 对比都依赖它）。"""
    text = "作\u2edb的\u2f00班\u2eda面"
    assert len(normalize_cjk(text)) == len(text)


def test_normalize_cjk_is_idempotent():
    once = normalize_cjk("\u2f00\u2eda作\u2edb")
    assert normalize_cjk(once) == once


def test_normalize_cjk_keeps_fullwidth_punctuation():
    """全角标点有语义，不能被顺手转成半角。"""
    text = "（一）：评分项目｜满分、实得分。"
    assert normalize_cjk(text) == text


def test_extract_text_normalizes_radicals(tmp_path):
    """端到端：走 extract_text() 也归一化。"""
    src = tmp_path / "a.txt"
    src.write_text("作\u2edb 与 \u2eda面", encoding="utf-8")
    assert extract_text(src) == "作风 与 页面"


# ---------- 部首残留兜底检测（D-43）----------


def test_radical_residue_is_silent_when_clean():
    assert check_radical_residue("作风与页面，全角标点（一）：评分项目。") is None
    assert check_radical_residue("") is None


def test_radical_residue_points_at_the_survivors():
    """U+2E80 既不在小表里、NFKC 也不管 —— 必须点名，不能静默通过。"""
    warning = check_radical_residue("详\u2e80 与 \u2e80")
    assert warning is not None
    assert "2 个" in warning
    assert "\u2e80" in warning


def test_radical_residue_summarizes_many_kinds():
    warning = check_radical_residue("".join(chr(c) for c in range(0x2E80, 0x2E87)))
    assert "7 个" in warning
    assert "等 7 种" in warning


def test_radical_residue_after_normalization_is_clean():
    """出口归一化过的文本，检测不该再抱怨（两道工序对得上）。"""
    assert check_radical_residue(normalize_cjk("作\u2edb \u2ec5 页\u2eda面")) is None
