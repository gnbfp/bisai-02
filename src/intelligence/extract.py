"""M1 前置步骤：文件 → 纯文本（**不涉及 LLM**）。

依据：requirements.md §7.5、docs/ARCHITECTURE.md §6.3、D-15 / D-21 / D-25。

产物形态就是一个 ``{ text: string }``：正文 + 表格区块。表格用
``find_tables()`` 单独抽，与正文用 ``=== 表格区域 ===`` 分隔；每行
``- 列名：值｜列名：值``（行间换行，避免跨行粘连 —— 待定义-33）。

职责边界：
  * 本模块**不调 LLM**、不落盘、不判定评分点。
  * ``check_weight_sum()`` / ``check_deadline()`` 的入参是 **M1 解析之后**的产物
    （``RubricPoint`` / ``AssignmentMeta``）。它们按 §6.3 的管线顺序放在这里，
    属于软校验：命中只返回一句警告，**绝不拒收**。
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Sequence

from src.models import AssignmentMeta, RubricPoint

__all__ = [
    "ExtractError",
    "TABLE_MARKER",
    "CELL_SEP",
    "VALUE_SEP",
    "ROW_PREFIX",
    "extract_text",
    "check_weight_sum",
    "check_deadline",
    "check_meta_fields",
    "normalize_cjk",
    "check_radical_residue",
]

TABLE_MARKER = "=== 表格区域 ==="
CELL_SEP = "｜"           # 列内分隔（待定义-33）
VALUE_SEP = "："          # 列名：值
ROW_PREFIX = "- "         # 每行前缀（待定义-33）

_TEXT_SUFFIXES = {"", ".txt", ".md", ".markdown", ".text"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".heic"}

WEIGHT_TARGET = 100.0     # 百分制满分
WEIGHT_TOLERANCE = 10.0   # 加总落在 [90, 110] 视为正常（§7.5）
DEADLINE_MIN_YEAR = 2000  # 早于此年份不可能真是截止时间，必是占位值（D-49）
# "能归一化到百分制"的落地口径：总分太小的量纲（如 20 分制）直接跳过，不误报。
# 之所以需要这道门：丢了几行的百分制评分表（如加总 85）与 20 分制在数值上无法
# 仅凭加总区分；低于本门限一律视为另一种量纲。此取值为工程默认，待真实作业书校准。
PERCENT_SCALE_MIN = 50.0


# ---------- CJK 字符归一化（D-43）----------

# CJK 部首与兼容汉字区。实测这两个区内 676 个字符有 NFKC 映射，
# 且**全部映射成单个汉字**（所以逐字符处理长度不变）。
_RADICAL_RANGES = ((0x2E80, 0x2FDF), (0xF900, 0xFAFF))

# NFKC 覆盖不到的简体部首（CJK Radicals Supplement）。
# 这两个区共 136 个字符无 NFKC 映射；扫过 7 份真实作业书（PDF + DOCX，2026-09-12），
# 命中过的只有下面 6 个 —— 其余漏网的由 check_radical_residue() 点名，不用硬背。
_RADICAL_FALLBACK = {
    "\u2ec5": "\u89c1",   # ⻅ -> 见（BIM 任务书 91 个部首里最后剩的 2 个）
    "\u2ec6": "\u89d2",   # ⻆ -> 角
    "\u2ed4": "\u95e8",   # ⻔ -> 门
    "\u2eda": "\u9875",   # ⻚ -> 页
    "\u2edb": "\u98ce",   # ⻛ -> 风
    "\u2ee2": "\u9a6c",   # ⻢ -> 马
}


def normalize_cjk(text: str) -> str:
    """把 PDF 抽出的部首字符还原成正常汉字（D-43）。

    字体缺 ToUnicode 映射时，PyMuPDF 会把常用汉字映射到康熙部首（一 -> U+2F00）
    或简体部首（页 -> U+2EDA）。实测真实课程任务书：2360 字里 176 个部首字符，
    正常「一」0 次。不还原的话，M1 的 quote 原文子串硬校验会把 LLM 写出的正常
    汉字引用（"页面"）判成幻觉，导致 M1 重试甚至整体降级。

    **逐字符处理，长度不变**；**只碰部首与兼容区，不动全角标点**。
    """
    if not text:
        return text
    return "".join(_normalize_char(ch) for ch in text)


def _normalize_char(ch: str) -> str:
    mapped = _RADICAL_FALLBACK.get(ch)
    if mapped is not None:
        return mapped
    if _is_radical_char(ch):
        return unicodedata.normalize("NFKC", ch)
    return ch


def _is_radical_char(ch: str) -> bool:
    cp = ord(ch)
    return any(low <= cp <= high for low, high in _RADICAL_RANGES)


# ---------- 部首残留兜底检测（D-43）----------


def check_radical_residue(text: str) -> str | None:
    """归一化之后还有没有漏网的部首字符。返回警告文案；``None`` = 干净。

    ``normalize_cjk()`` 只认 NFKC + 一张小表，U+2E80–U+2FDF 里那 130 个两边都兜不住
    的字符**不会报错**，只会安静地污染 M1 的 ``quote`` 校验（D-43 讲的危害）。
    所以在出口之后数一遍，让漏网的以软警告的形式点名露面 —— 新部首从此不会静默通过，
    由人决定要不要往 ``_RADICAL_FALLBACK`` 里补一行。

    与 ``check_weight_sum()`` 同款：**软警告不拒收**，调用方把返回值拼进回复即可。
    """
    residue = [ch for ch in (text or "") if _is_radical_char(ch)]
    if not residue:
        return None
    kinds = sorted(set(residue))
    shown = " ".join(kinds[:5])
    extra = f" 等 {len(kinds)} 种" if len(kinds) > 5 else ""
    return f"文本里有 {len(residue)} 个部首字符未归一化，请核对原文（如 {shown}{extra}）"


class ExtractError(RuntimeError):
    """文件抽不出文本 —— 扫描版 PDF / 图片 / 不认识的格式。消息直接可以发给用户。"""


def extract_text(path: Path | str) -> str:
    """按后缀分派：PDF / Word / 纯文本，最后统一做 CJK 归一化（D-43）。

    归一化放在唯一出口，任何新格式进来都自动受益 —— 别在三个分支里各写一遍。
    """
    return normalize_cjk(_extract_raw(Path(path)))


def _extract_raw(p: Path) -> str:
    """按后缀分派：PDF / Word / 纯文本。图片与扫描版直接拒收（§7.5）。"""
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _pdf_to_text(p)
    if suffix == ".docx":
        return _docx_to_text(p)
    if suffix in _TEXT_SUFFIXES:
        return p.read_text(encoding="utf-8")
    if suffix in _IMAGE_SUFFIXES:
        raise ExtractError(f"图片不能抽取文字（OCR 已砍，P2）：{p.name}。请发文字版作业书。")
    raise ExtractError(f"不认识的文件类型 {suffix or '(无后缀)'}：{p.name}。支持 PDF / DOCX / TXT。")


# ---------- PDF ----------


def _pdf_to_text(path: Path) -> str:
    import fitz  # PyMuPDF：懒导入，纯文本链路不因它缺失而 import 失败

    body_chunks: list[str] = []
    table_lines: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            body_chunks.append(page.get_text("text"))
            for table in page.find_tables().tables:
                names, data = _split_header(table)
                table_lines.extend(_rows_to_lines(data, names))

    body = "\n".join(body_chunks).strip()
    if not body and not table_lines:
        raise ExtractError(
            f"PDF 没有文字层（扫描版或图片版）：{path.name}。请发文字版作业书。"
        )
    return _assemble(body, table_lines)


def _split_header(table) -> tuple[list[str] | None, list[list[str]]]:
    """拆出列名与数据行。PyMuPDF 的 ``header.external`` 决定表头是否已在 extract() 里。"""
    rows = [[_cell(c) for c in row] for row in table.extract()]
    rows = [row for row in rows if any(row)]
    if not rows:
        return None, []
    try:
        names = [_cell(n) for n in table.header.names]
        external = bool(table.header.external)
    except Exception:                      # 低版本 / 未识别表头：退回"首行即表头"
        names, external = None, False
    if not names or not any(names):
        return rows[0], rows[1:]
    return names, rows if external else rows[1:]


# ---------- Word ----------


def _docx_to_text(path: Path) -> str:
    from docx import Document  # python-docx

    doc = Document(path)
    body = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    table_lines: list[str] = []
    for table in doc.tables:
        rows = [[_cell(c.text) for c in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if not rows:
            continue
        table_lines.extend(_rows_to_lines(rows[1:], rows[0]))

    if not body and not table_lines:
        raise ExtractError(f"Word 文档没有可抽取的文字：{path.name}。")
    return _assemble(body, table_lines)


# ---------- 共享格式 ----------


def _cell(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())      # 压掉单元格里的换行与连续空白


def _rows_to_lines(rows: Sequence[Sequence[str]], names: Sequence[str] | None) -> list[str]:
    """数据行 → ``- 列名：值｜列名：值``；空单元格跳过。"""
    lines: list[str] = []
    for row in rows:
        pairs = []
        for index, value in enumerate(row):
            if not value:
                continue
            key = names[index] if names and index < len(names) and names[index] else f"列{index + 1}"
            pairs.append(f"{key}{VALUE_SEP}{value}")
        if pairs:
            lines.append(ROW_PREFIX + CELL_SEP.join(pairs))
    return lines


def _assemble(body: str, table_lines: Sequence[str]) -> str:
    parts = [body] if body else []
    if table_lines:
        parts.append(TABLE_MARKER)
        parts.extend(table_lines)
    return "\n".join(parts)


# ---------- 权重加总软校验（§7.5，原待定义-32 / D-25）----------


def check_weight_sum(points: Sequence[RubricPoint]) -> str | None:
    """权重加总软校验。返回警告文案；``None`` = 跳过或正常。

    分支（§7.5 / D-39 / D-48）：
      * 全部没有 ``weight`` → 整体跳过（§6.1 允许无分值的评分点，报警会成噪音）；
      * **部分有、部分没有** → 软警告"疑似把正文要求当成了评分点"：真评分点通常整节
        都标分值，一半带分一半不带，往往是 M1 把正文要求也收进来了；
      * 全都有 → 量纲是百分制（``>= PERCENT_SCALE_MIN``）时校验加总是否接近 100；
        20 分制之类整体跳过，不误报。

    **软警告不拒收**：调用方（M1）把返回值拼进给用户的回复即可。
    """
    if not points:
        return None
    weights = [p.weight for p in points]
    if any(weight is None for weight in weights):
        if any(weight is not None for weight in weights):
            unweighted = sum(1 for weight in weights if weight is None)
            return (
                f"评分点里有 {unweighted}/{len(points)} 条没有分值，"
                f"疑似把正文要求当成了评分点，请对照原文核对"
            )
        return None      # 全都没有分值 → 保持现状（§6.1 明确允许无分值的评分点）
    total = float(sum(weights))              # type: ignore[arg-type]
    if total <= 0 or total < PERCENT_SCALE_MIN:
        return None
    low, high = WEIGHT_TARGET - WEIGHT_TOLERANCE, WEIGHT_TARGET + WEIGHT_TOLERANCE
    if low <= total <= high:
        return None
    return (
        f"评分标准解析疑似丢行，请对照原文核对"
        f"（评分点权重加总 = {total:g}，应接近 {WEIGHT_TARGET:g}）"
    )


# ---------- 截止时间软校验（D-49）----------


def check_deadline(meta: AssignmentMeta | None) -> str | None:
    """截止时间软校验。返回警告文案；``None`` = 正常（D-49）。

    schema 放开 ``deadline`` 之后，提示词"没有就留空"才真的成立：空值不是错误，
    但要**明确报出来**，否则报告上是一个肉眼看不出的空字符串。同时挡住 LLM 为了
    跳出校验而死填的占位值（实测出现过 ``1970-01-01T00:00``）。

    门限取"年份 < 2000"而不是"早于今天"—— 真实但已过期的截止时间不该被误报。
    与 ``check_weight_sum()`` 同款：**软警告、不拒收**。
    """
    raw = (getattr(meta, "deadline", "") or "").strip()
    if not raw:
        return "作业书里没读到明确的截止时间，报告按「未标注」显示"
    match = re.match(r"^(\d{4})-", raw)
    if not match or int(match.group(1)) < DEADLINE_MIN_YEAR:
        return f"截止时间 {raw!r} 很可能是占位值（原文没有明确日期），请对照原文核对"
    return None


def check_meta_fields(meta: AssignmentMeta | None) -> str | None:
    """元信息缺件软校验（2026-09-17 PM 拍 A）。返回警告文案；``None`` = 齐。

    schema 放开 ``course / title / submission`` 之后，「只有格式要求、没有课程名」这类
    文档（U3 的典型输入）不再卡在 M1 校验里重试到 ``LLMError``；但空值必须**明确
    报出来**，否则清单抬头就是一个肉眼看不出的空档 —— 与 ``check_deadline()`` 同款：
    **软警告、不拒收**。``source_file`` 不在此列：它是程序给的事实，仍然必填。
    """
    labels = (("course", "课程名"), ("title", "作业标题"), ("submission", "提交物"))
    missing = [label for name, label in labels if not (getattr(meta, name, "") or "").strip()]
    if not missing:
        return None
    return f"作业书里没读到{'/'.join(missing)}，报告按「未标注」显示"
